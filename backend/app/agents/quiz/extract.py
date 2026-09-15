import io
import re
from pathlib import Path

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".docx", ".txt", ".md", ".vtt", ".srt"}

_TIMESTAMP = re.compile(r"\d{1,2}:\d{2}(:\d{2})?[.,]\d{3}\s*-->\s*\d{1,2}:\d{2}(:\d{2})?[.,]\d{3}.*")


def _clean_transcript(text: str) -> str:
    lines = []
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped or stripped == "WEBVTT" or stripped.isdigit() or _TIMESTAMP.match(stripped):
            continue
        lines.append(re.sub(r"<[^>]+>", "", stripped))
    return " ".join(lines)


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename or "").suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)

    if ext == ".docx":
        from docx import Document
        document = Document(io.BytesIO(data))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)

    if ext in {".txt", ".md", ".vtt", ".srt"}:
        text = data.decode("utf-8", errors="ignore")
        return _clean_transcript(text) if ext in {".vtt", ".srt"} else text

    raise ValueError(
        f"Unsupported file type '{ext or 'unknown'}'. Upload PDF, PPTX, DOCX, TXT, MD, or a video transcript (VTT/SRT)."
    )
